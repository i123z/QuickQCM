import json

from pylatex import Document, Section, Itemize, Enumerate, Description, Command, NoEscape
from pylatex import Document, Section, Subsection, Command
from pylatex.utils import italic, NoEscape
from pylatex.package import Package
from pylatex import Document, Section, UnsafeCommand
from pylatex.base_classes import Environment, CommandBase, Arguments
import os
import sys
import csv
import openpyxl
from openpyxl import workbook



with open("LatexQCM\\" + sys.argv[1]) as f:
    data = json.load(f)


print(data["niveau"])
dossier_fichiers_excel = 'LatexQCM\\Listes etudiants excel\\'
nom_fichier_excel_emails = data['niveau'] + ' - Email.xlsx'
wb_emails = openpyxl.load_workbook(dossier_fichiers_excel + nom_fichier_excel_emails)
feuille_emails = wb_emails['JM']
indice_debut = 7
indice_fin = feuille_emails.max_row + 1
indice_nom = 'B'
indice_prenom = 'C'
indice_id = 'E'



def genererQCM(data):

    doc = Document(documentclass = 'book')
    nombreQuestions = len(data['questions'])
    doc.packages.append(NoEscape('\\input{structure}'))
    
    for j in range(indice_debut, indice_fin):
        ########################################################
        doc.append(NoEscape("\\newpage" ))
        doc.append(NoEscape("\\thispagestyle{empty}"))
        doc.append(NoEscape("\\vskip-40mm	\\includegraphics[scale=0.5]{logo.png} \\\\"))
        doc.append(NoEscape(" \\begin{flushright}  \\vskip-20mm   Professeur: " + data["prof"] + "\\vskip15mm  \\end{flushright}"))
        doc.append(NoEscape(data["niveau"]))
        doc.append(NoEscape(" \\begin{flushright}  \\vskip-7mm"+    data["date"] + " \\end{flushright}"))
        doc.append(NoEscape("\\begin{center}   \\begin{Large}" + data["type_qcm"] + " : "+ data["matiere"].upper() + "\\end{Large} \\end{center}"))
        doc.append(NoEscape("Durée : " + data["duree"]+ "h" ))


        doc.append(NoEscape(" \\begin{center} { \\large CONSIGNES SPÉCIFIQUES } \\\\\\ Lisez soigneusement les consignes ci-dessous afin de réussir au mieux cette épreuve : \\end{center} "))
        with doc.create(Itemize()) as item:
            if data["calculatrice"]==0:
                item.add_item("L'usage de la calculatrice ou de tout autre appareil électronique est interdit.")
            else:
                item.add_item("L'usage de la calculatrice ou de tout autre appareil électronique est autorisé.")
            if data["documents"]==0:
                item.add_item("Aucun autre document que ce sujet et sa grille réponse n'est autorisé.")
            else:
                item.add_item("Documents autorisés.")
            if data["autreConsignes"]!="":
                for i in data["autreConsignes"].split("\n"):
                    item.add_item(i)
                    
            item.add_item(NoEscape("Pour chacune des questions, indiquez sur la feuille de réponses ci-jointes, si les affirmations A, B, C et D sont (\\textbf{V}) vraies ou (\\textbf{F}) fausses en faisant une croix dans la colonne correspondant à votre choix. Vous ne pouvez pas faire de ratures. En cas d'erreur, utilisez la deuxième colonne de réponse. Si la deuxième colonne comporte au moins une réponse, la première colonne ne sera pas corrigée, c'est la deuxième qui sera prise en considération."))
            item.add_item(NoEscape("Chaque réponse exacte est gratifiée de "+ str(data["notation"][0]) +" points, tandis que chaque réponse fausse est pénalisée par "+ str(data["notation"][1]) +" point. \\\\ 	Parmi les quatre propositions de chacune des questions \\textbf{de 1 à "+ str(nombreQuestions)+ "}, une seule est vraie, les autres sont fausses. \\\\ 	Par exemple : Pour indiquer que l'affirmation $B$ est Vraie, cocher les cases comme suit:  \\\\ \\begin{center}	\\includegraphics[scale=0.8]{reponses.png} \\end{center}"))

            for i in range(nombreQuestions):
                doc.append(NoEscape("\\thispagestyle{empty}"))
                doc.append(NoEscape("\\begin{exercise}"))
                doc.append(NoEscape("\\textbf{" + data["questions"][i]["texte"] +  " }"))
                with doc.create(Enumerate(enumeration_symbol=r"\textbf{\Alph*. }"), ) as enum:
                    enum.add_item(NoEscape(data["questions"][i]["options"][0]))
                    enum.add_item(NoEscape(data["questions"][i]["options"][1]))
                    enum.add_item(NoEscape(data["questions"][i]["options"][2]))
                    enum.add_item(NoEscape(data["questions"][i]["options"][3]))
                doc.append(NoEscape("\\end{exercise}"))
        doc.append(NoEscape("\\newpage"))
        doc.append(NoEscape("\\thispagestyle{empty}"))
        nom = str(feuille_emails[indice_nom + str(j)].value)
        prenom = str(feuille_emails[indice_prenom + str(j)].value)
        identidfiant =  str(feuille_emails[indice_id + str(j)].value)
        doc.append(NoEscape(data["type_qcm"] + " : " + data["matiere"].upper() + " $\\qquad \\qquad \\qquad \\qquad \\qquad \\qquad \\qquad \\qquad$ " + nom + ' ' + prenom))
        doc.append(NoEscape("\\begin{flushright}"))
        doc.append(NoEscape("\\begin{tabular}{|l|}"))
        doc.append(NoEscape("\\hline"))
        doc.append(NoEscape(" \\\\"))
        identidfiant = str(identidfiant)
        text = ""
        for c in identidfiant:
            text += c + "~"
        doc.append(NoEscape("\\thispagestyle{empty}"))
        doc.append(NoEscape("Identifiant: $\\quad$ {\\Large " + text + "}"))
        doc.append(NoEscape(" \\\\"))
        doc.append(NoEscape("\\hline"))
        doc.append(NoEscape("\\end{tabular}"))
        doc.append(NoEscape("\\end{flushright}"))
        doc.append(NoEscape("\\begin{center}"))
        for i in range(nombreQuestions):
            doc.append(NoEscape("\\begin{tabular}{| l l l l l |}"))
            doc.append(NoEscape("\\hline"))
            doc.append(NoEscape(" & & & & \\\\"))
            doc.append(NoEscape("Question " + str(i+1) + "\\qquad \\qquad\\ & & & & \\\\"))
            doc.append(NoEscape(" & A $\\qquad \\square \\qquad$ & B $\\qquad \\square \\qquad$ & C $\\qquad \\square \\qquad$ & D $\\qquad \\square \\qquad$ \\\\ "))
            doc.append(NoEscape(" & & & &  \\\\"))
            doc.append(NoEscape("\\hline"))
            doc.append(NoEscape(" & & & &  \\\\"))
            doc.append(NoEscape("Choix 2 & A $\\qquad \\square \\qquad$ & B $\\qquad \\square \\qquad$ & C $\\qquad \\square \\qquad$ & D $\\qquad \\square \\qquad$ \\\\ "))
            doc.append(NoEscape(" & & & &  \\\\"))
            doc.append(NoEscape("\\hline"))
            doc.append(NoEscape("\\end{tabular}"))
            doc.append(NoEscape("\\\\ \\vskip3mm"))
            doc.append(NoEscape("\\thispagestyle{empty}"))
        doc.append(NoEscape("\\end{center}"))
        #######################################################################

    doc.generate_pdf("LatexQCM\\" + data["nomFichier"], clean_tex=False, compiler='pdfLaTeX')


genererQCM(data)